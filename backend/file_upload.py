"""File upload handling for research documents (PDF, CSV, TXT, DOCX)."""

import os
import uuid
import json
from typing import Optional
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from database import get_db, User, Project
from auth import get_current_user
from config import settings

router = APIRouter(prefix="/api/files", tags=["files"])

UPLOAD_DIR = Path("/app/uploads")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {
    ".pdf", ".csv", ".txt", ".docx", ".doc", ".xlsx", ".md", ".tex", ".bib",
    ".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp",
    ".pptx", ".ppt",
}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp"}
SLIDE_EXTENSIONS = {".pptx", ".ppt"}


def _extract_text_from_file(file_path: Path, content_type: str) -> str:
    """Extract text content from uploaded files for context injection."""
    suffix = file_path.suffix.lower()

    if suffix == ".txt" or suffix == ".md" or suffix == ".tex" or suffix == ".bib":
        return file_path.read_text(errors="replace")[:50000]

    if suffix == ".csv":
        text = file_path.read_text(errors="replace")
        lines = text.split("\n")
        # Return header + first 200 rows
        return "\n".join(lines[:201])

    if suffix == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(file_path))
            text = ""
            for page in doc:
                text += page.get_text()
                if len(text) > 50000:
                    break
            doc.close()
            return text[:50000]
        except ImportError:
            return "[PDF content - install PyMuPDF for text extraction]"
        except Exception:
            return "[Could not extract PDF text]"

    if suffix in (".docx", ".doc"):
        try:
            import docx
            doc = docx.Document(str(file_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            return text[:50000]
        except ImportError:
            return "[DOCX content - install python-docx for text extraction]"
        except Exception:
            return "[Could not extract document text]"

    if suffix in IMAGE_EXTENSIONS:
        return f"[Image: {file_path.name} — describe what you see in this image if relevant to the research]"

    if suffix in SLIDE_EXTENSIONS:
        try:
            from pptx import Presentation
            prs = Presentation(str(file_path))
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        texts.append(shape.text)
            return "\n".join(texts)[:50000]
        except ImportError:
            return "[PPTX content - install python-pptx for text extraction]"
        except Exception:
            return "[Could not extract slide text]"

    return f"[File: {file_path.name}]"


@router.post("/chat-upload")
async def chat_upload(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
):
    """Upload a file inline during chat (not tied to a project). Returns extracted text."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type {ext} not allowed. Supported: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File too large. Maximum 10MB.")

    file_id = str(uuid.uuid4())[:8]
    safe_name = f"{file_id}_{file.filename.replace(' ', '_')}"
    user_dir = UPLOAD_DIR / str(user.id) / "chat"
    user_dir.mkdir(parents=True, exist_ok=True)
    file_path = user_dir / safe_name
    file_path.write_bytes(content)

    is_image = ext in IMAGE_EXTENSIONS
    extracted_text = "" if is_image else _extract_text_from_file(file_path, file.content_type or "")

    return {
        "id": file_id,
        "name": file.filename,
        "size": len(content),
        "type": ext,
        "is_image": is_image,
        "extracted_text": extracted_text[:20000],
        "path": str(file_path),
    }


@router.post("/upload")
async def upload_file(
    file: UploadFile = File(...),
    project_id: int = Form(...),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    # Check tier allows file upload
    tier_info = settings.tier_limits.get(user.subscription_tier.value, {})
    if not tier_info.get("file_upload", False):
        raise HTTPException(
            status_code=403,
            detail="File upload requires Pro or Lab Group plan.",
        )

    # Validate project ownership
    result = await db.execute(
        select(Project).where(Project.id == project_id, Project.user_id == user.id)
    )
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    # Validate file
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type {ext} not allowed. Supported: {', '.join(ALLOWED_EXTENSIONS)}",
        )

    # Read and check size
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File too large. Maximum 10MB.")

    # Save file
    file_id = str(uuid.uuid4())[:8]
    safe_name = f"{file_id}_{file.filename.replace(' ', '_')}"
    user_dir = UPLOAD_DIR / str(user.id)
    user_dir.mkdir(parents=True, exist_ok=True)
    file_path = user_dir / safe_name
    file_path.write_bytes(content)

    # Extract text for context injection
    extracted_text = _extract_text_from_file(file_path, file.content_type or "")

    # Update project context_files_json
    files_list = project.context_files_json or []
    file_entry = {
        "id": file_id,
        "name": file.filename,
        "stored_name": safe_name,
        "size": len(content),
        "type": ext,
        "extracted_text_preview": extracted_text[:500],
    }
    files_list.append(file_entry)

    await db.execute(
        update(Project)
        .where(Project.id == project_id)
        .values(context_files_json=files_list)
    )
    await db.commit()

    return {
        "id": file_id,
        "name": file.filename,
        "size": len(content),
        "type": ext,
        "extracted_length": len(extracted_text),
        "project_id": project_id,
    }


@router.get("/project/{project_id}")
async def list_project_files(
    project_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Project).where(Project.id == project_id, Project.user_id == user.id)
    )
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    return project.context_files_json or []


@router.delete("/{file_id}")
async def delete_file(
    file_id: str,
    project_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Project).where(Project.id == project_id, Project.user_id == user.id)
    )
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")

    files_list = project.context_files_json or []
    file_entry = next((f for f in files_list if f["id"] == file_id), None)
    if not file_entry:
        raise HTTPException(status_code=404, detail="File not found")

    # Delete from disk
    file_path = UPLOAD_DIR / str(user.id) / file_entry["stored_name"]
    if file_path.exists():
        file_path.unlink()

    # Remove from project
    files_list = [f for f in files_list if f["id"] != file_id]
    await db.execute(
        update(Project)
        .where(Project.id == project_id)
        .values(context_files_json=files_list)
    )
    await db.commit()

    return {"message": "File deleted"}


def get_project_context_text(project: Project) -> str:
    """Build context string from project files for LLM injection."""
    if not project or not project.context_files_json:
        return ""

    context_parts = []
    for f in project.context_files_json:
        file_path = UPLOAD_DIR / str(project.user_id) / f["stored_name"]
        if file_path.exists():
            text = _extract_text_from_file(file_path, "")
            if text:
                context_parts.append(f"--- File: {f['name']} ---\n{text[:10000]}")

    if context_parts:
        return "\n\nReference documents:\n" + "\n\n".join(context_parts)
    return ""
